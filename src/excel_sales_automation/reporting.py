"""Formatted Excel report generation."""

from __future__ import annotations

import importlib
import importlib.util
import logging
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from excel_sales_automation import __version__
from excel_sales_automation.models import ComparisonResult

openpyxl = (
    importlib.import_module("openpyxl")
    if importlib.util.find_spec("openpyxl") is not None
    else None
)

logger = logging.getLogger(__name__)

REQUIRED_SHEETS = [
    "Summary",
    "Matched",
    "Missing_In_Baseline",
    "Missing_In_Comparison",
    "Duplicates",
    "Invalid_Rows",
]
FORMULA_DEFINITIONS = [
    "price_difference = comparison_price - baseline_price",
    "price_difference_rate = price_difference / baseline_price",
    "price_difference_over_comparison = price_difference / comparison_price",
    "gross_margin = (comparison_price - cost) / comparison_price",
    "auto margin formula uses gross_margin when cost exists; "
    "otherwise price_difference_over_comparison",
]
ASSUMPTIONS = [
    "Product rows are matched by normalized product_code.",
    "Duplicate product codes are reported and excluded from clean matched calculations "
    "when policy is report.",
    "Division by zero and missing numeric inputs create calculation warnings instead of crashing.",
    "Output is local-only and no scraping, browser automation, or external services are used.",
]


class ReportExportError(ValueError):
    """Raised when a report workbook cannot be exported."""


def validate_report_output_path(path: str | Path) -> Path:
    """Validate the destination path for a generated Excel report."""
    output_path = Path(path)
    if output_path.suffix.lower() != ".xlsx":
        raise ReportExportError("Report output path must use the .xlsx extension.")
    return output_path


def export_excel_report(
    result: ComparisonResult,
    output_path: str | Path,
    *,
    input_file_names: list[str] | None = None,
    config_path: str | Path | None = None,
    generated_at: datetime | None = None,
) -> Path:
    """Export a formatted multi-sheet Excel report."""
    report_path = validate_report_output_path(output_path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    timestamp = generated_at or datetime.now(UTC)
    logger.info("Exporting Excel report to %s", report_path)

    sheets = _build_report_sheets(
        result,
        input_file_names=input_file_names or [],
        config_path=config_path,
        generated_at=timestamp,
    )

    try:
        if openpyxl is not None:
            _write_with_openpyxl(sheets, report_path)
        else:
            _write_minimal_xlsx(sheets, report_path)
    except OSError as exc:
        raise ReportExportError(f"Could not write report workbook: {report_path}: {exc}") from exc

    return report_path


def table_to_records(table: Any) -> list[dict[str, Any]]:
    """Convert supported table-like objects to records."""
    if table is None:
        return []
    if isinstance(table, list):
        return [dict(record) for record in table]
    if hasattr(table, "to_dict"):
        return [dict(record) for record in table.to_dict("records")]
    if hasattr(table, "records"):
        return [dict(record) for record in table.records]
    raise ReportExportError(f"Unsupported report table type: {type(table)!r}")


def _build_report_sheets(
    result: ComparisonResult,
    *,
    input_file_names: list[str],
    config_path: str | Path | None,
    generated_at: datetime,
) -> dict[str, list[dict[str, Any]]]:
    matched_records = table_to_records(result.matched_rows)
    sheets = {
        "Summary": _summary_records(
            result, matched_records, input_file_names, config_path, generated_at
        ),
        "Matched": matched_records,
        "Missing_In_Baseline": table_to_records(result.missing_in_baseline),
        "Missing_In_Comparison": table_to_records(result.missing_in_comparison),
        "Duplicates": table_to_records(result.duplicates),
        "Invalid_Rows": table_to_records(result.invalid_rows),
    }
    return sheets


def _summary_records(
    result: ComparisonResult,
    matched_records: list[dict[str, Any]],
    input_file_names: list[str],
    config_path: str | Path | None,
    generated_at: datetime,
) -> list[dict[str, Any]]:
    summary = [
        {"Section": "Run Metadata", "Metric": "generated_at", "Value": generated_at.isoformat()}
    ]
    summary.append({"Section": "Run Metadata", "Metric": "report_version", "Value": __version__})
    if config_path is not None:
        summary.append(
            {"Section": "Run Metadata", "Metric": "config_path", "Value": str(config_path)}
        )
    for index, file_name in enumerate(input_file_names, start=1):
        summary.append(
            {"Section": "Run Metadata", "Metric": f"input_file_{index}", "Value": file_name}
        )

    for metric, value in result.summary.items():
        summary.append({"Section": "Counts", "Metric": metric, "Value": value})

    status_counts = _calculation_status_counts(matched_records)
    for metric, value in status_counts.items():
        summary.append({"Section": "Calculation Status", "Metric": metric, "Value": value})

    for formula in FORMULA_DEFINITIONS:
        summary.append({"Section": "Formula", "Metric": "definition", "Value": formula})
    for assumption in ASSUMPTIONS:
        summary.append({"Section": "Assumption", "Metric": "note", "Value": assumption})
    return summary


def _calculation_status_counts(records: list[dict[str, Any]]) -> dict[str, int]:
    counts = {
        "calculation_ok_count": 0,
        "calculation_warning_count": 0,
        "calculation_invalid_count": 0,
    }
    for record in records:
        status = record.get("calculation_status")
        if status == "ok":
            counts["calculation_ok_count"] += 1
        elif status == "warning":
            counts["calculation_warning_count"] += 1
        elif status == "invalid":
            counts["calculation_invalid_count"] += 1
    return counts


def _write_with_openpyxl(sheets: dict[str, list[dict[str, Any]]], output_path: Path) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for sheet_name, records in sheets.items():
        worksheet = workbook.create_sheet(sheet_name)
        _write_openpyxl_records(worksheet, records)
        if worksheet.max_row >= 1 and worksheet.max_column >= 1:
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for column_cells in worksheet.columns:
                letter = get_column_letter(column_cells[0].column)
                width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 50)
                worksheet.column_dimensions[letter].width = max(width, 12)
            _format_openpyxl_numeric_columns(worksheet)

    workbook.save(output_path)


def _write_openpyxl_records(worksheet: Any, records: list[dict[str, Any]]) -> None:
    if not records:
        worksheet.append(["No rows"])
        return
    headers = _headers(records)
    worksheet.append(headers)
    for record in records:
        worksheet.append([record.get(header) for header in headers])


def _format_openpyxl_numeric_columns(worksheet: Any) -> None:
    percent_columns = {"price_difference_rate", "margin_rate"}
    money_fragments = ("price", "cost", "difference")
    headers = [cell.value for cell in worksheet[1]]
    for column_index, header in enumerate(headers, start=1):
        if header in percent_columns:
            number_format = "0.00%"
        elif any(fragment in str(header) for fragment in money_fragments):
            number_format = "#,##0.00"
        else:
            continue
        for row in worksheet.iter_rows(min_row=2, min_col=column_index, max_col=column_index):
            row[0].number_format = number_format


def _write_minimal_xlsx(sheets: dict[str, list[dict[str, Any]]], output_path: Path) -> None:
    with zipfile.ZipFile(output_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", _content_types_xml(len(sheets)))
        archive.writestr("_rels/.rels", _root_rels_xml())
        archive.writestr("xl/workbook.xml", _workbook_xml(list(sheets)))
        archive.writestr("xl/_rels/workbook.xml.rels", _workbook_rels_xml(len(sheets)))
        archive.writestr("xl/styles.xml", _styles_xml())
        for index, records in enumerate(sheets.values(), start=1):
            archive.writestr(f"xl/worksheets/sheet{index}.xml", _worksheet_xml(records))


def _headers(records: list[dict[str, Any]]) -> list[str]:
    headers: list[str] = []
    for record in records:
        for key in record:
            if key not in headers:
                headers.append(key)
    return headers


def _worksheet_xml(records: list[dict[str, Any]]) -> str:
    rows: list[list[Any]]
    if records:
        headers = _headers(records)
        rows = [headers] + [[record.get(header, "") for header in headers] for record in records]
    else:
        rows = [["No rows"]]
    row_xml = "".join(_row_xml(row_index, row) for row_index, row in enumerate(rows, start=1))
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<sheetData>{row_xml}</sheetData><autoFilter ref="A1:Z1"/></worksheet>'
    )


def _row_xml(row_index: int, values: list[Any]) -> str:
    cells = "".join(
        _cell_xml(row_index, column_index, value)
        for column_index, value in enumerate(values, start=1)
    )
    return f'<row r="{row_index}">{cells}</row>'


def _cell_xml(row_index: int, column_index: int, value: Any) -> str:
    reference = f"{_column_letter(column_index)}{row_index}"
    if isinstance(value, int | float) and not isinstance(value, bool):
        return f'<c r="{reference}"><v>{value}</v></c>'
    return f'<c r="{reference}" t="inlineStr"><is><t>{escape(str(value))}</t></is></c>'


def _column_letter(index: int) -> str:
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _content_types_xml(sheet_count: int) -> str:
    sheet_overrides = "".join(
        (
            f'<Override PartName="/xl/worksheets/sheet{index}.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.worksheet+xml"/>'
        )
        for index in range(1, sheet_count + 1)
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" '
        'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/styles.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        f"{sheet_overrides}</Types>"
    )


def _root_rels_xml() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="xl/workbook.xml"/></Relationships>'
    )


def _workbook_xml(sheet_names: list[str]) -> str:
    sheets_xml = "".join(
        f'<sheet name="{escape(name)}" sheetId="{index}" r:id="rId{index}"/>'
        for index, name in enumerate(sheet_names, start=1)
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f"<sheets>{sheets_xml}</sheets></workbook>"
    )


def _workbook_rels_xml(sheet_count: int) -> str:
    sheet_rels = "".join(
        (
            f'<Relationship Id="rId{index}" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
            f'Target="worksheets/sheet{index}.xml"/>'
        )
        for index in range(1, sheet_count + 1)
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'{sheet_rels}<Relationship Id="rId{sheet_count + 1}" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" '
        'Target="styles.xml"/></Relationships>'
    )


def _styles_xml() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
        '<fills count="1"><fill><patternFill patternType="none"/></fill></fills>'
        '<borders count="1"><border/></borders>'
        '<cellStyleXfs count="1"><xf/></cellStyleXfs>'
        '<cellXfs count="1"><xf/></cellXfs></styleSheet>'
    )
